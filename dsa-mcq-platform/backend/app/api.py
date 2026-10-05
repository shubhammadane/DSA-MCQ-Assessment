from fastapi import APIRouter, Depends, HTTPException, status, Header, Query, UploadFile, File, Form
from sqlalchemy.orm import Session
from sqlalchemy import func, or_
from datetime import datetime, timedelta
from typing import List, Optional, Dict
import io
import csv
import re
import random
import openpyxl
from pypdf import PdfReader
from fastapi.responses import StreamingResponse

from app.database import get_db
from app import models, schemas
from app.config import settings
from app.auth import (
    create_access_token,
    get_current_admin,
    get_current_student,
    hash_password,
    verify_password
)
from app.questions_data import FIXED_QUESTIONS, QUESTIONS_BY_ID

router = APIRouter(prefix="/api")


# =========================================================================
# HELPER FUNCTIONS
# =========================================================================

def get_or_create_settings(db: Session) -> models.AssessmentSetting:
    setting = db.query(models.AssessmentSetting).first()
    if not setting:
        setting = models.AssessmentSetting(
            question_count=50,
            time_limit_minutes=60,
            updated_at=datetime.utcnow()
        )
        db.add(setting)
        db.commit()
        db.refresh(setting)
    return setting


def calculate_attempt_score(attempt: models.Attempt, db: Session):
    """Recalculate score and percentage for an attempt based on actual questions."""
    answers = db.query(models.StudentAnswer).filter(models.StudentAnswer.attempt_id == attempt.id).all()
    correct_count = sum(1 for a in answers if a.is_correct)
    wrong_count = sum(1 for a in answers if not a.is_correct and a.selected_answer is not None)
    total_q = attempt.total_questions or len(answers) or 1
    
    attempt.correct_answers = correct_count
    attempt.wrong_answers = wrong_count
    attempt.score = correct_count
    attempt.percentage = round((correct_count / total_q) * 100.0, 2)


# =========================================================================
# 1. ADMIN AUTHENTICATION
# =========================================================================

@router.post("/auth/login", response_model=schemas.Token)
def login(credentials: schemas.AdminLogin):
    if credentials.username == settings.ADMIN_USERNAME and credentials.password == settings.ADMIN_PASSWORD:
        access_token = create_access_token(data={"sub": credentials.username, "role": "admin"})
        return {"access_token": access_token, "token_type": "bearer"}
    raise HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Incorrect admin username or password"
    )


# =========================================================================
# 2. ACADEMIC STRUCTURE & DEPARTMENTS
# =========================================================================

@router.get("/academic-structure", response_model=schemas.AcademicStructureResponse)
def get_academic_structure(db: Session = Depends(get_db)):
    depts = db.query(models.Department).filter(models.Department.is_active == True).order_by(models.Department.name.asc()).all()
    progs = db.query(models.Program).filter(models.Program.is_active == True).order_by(models.Program.id.asc()).all()
    years = db.query(models.AcademicYear).filter(models.AcademicYear.is_active == True).order_by(models.AcademicYear.program_id.asc(), models.AcademicYear.year_number.asc()).all()
    sems = db.query(models.Semester).filter(models.Semester.is_active == True).order_by(models.Semester.program_id.asc(), models.Semester.semester_number.asc()).all()
    return {
        "departments": depts,
        "programs": progs,
        "years": years,
        "semesters": sems
    }


@router.get("/departments", response_model=List[schemas.DepartmentOut])
def get_departments(all_status: bool = False, db: Session = Depends(get_db)):
    q = db.query(models.Department)
    if not all_status:
        q = q.filter(models.Department.is_active == True)
    return q.order_by(models.Department.name.asc()).all()


@router.post("/departments", response_model=schemas.DepartmentOut, status_code=status.HTTP_201_CREATED)
def create_department(dept_in: schemas.DepartmentCreate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    name = dept_in.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Department name is required.")
    existing = db.query(models.Department).filter(models.Department.name.ilike(name)).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Department '{name}' already exists.")
    dept = models.Department(
        name=name,
        code=dept_in.code.strip() if dept_in.code else name[:6].upper().replace(" ", ""),
        is_active=True,
        created_at=datetime.utcnow()
    )
    db.add(dept)
    db.commit()
    db.refresh(dept)
    return dept


@router.patch("/departments/{department_id}/status", response_model=schemas.DepartmentOut)
def toggle_department_status(department_id: int, status_in: schemas.QuestionStatusUpdate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    dept = db.query(models.Department).filter(models.Department.id == department_id).first()
    if not dept:
        raise HTTPException(status_code=404, detail="Department not found.")
    dept.is_active = status_in.is_active
    db.commit()
    db.refresh(dept)
    return dept


# =========================================================================
# 3. SUBJECT MANAGEMENT (MANUALLY ADDED ONLY)
# =========================================================================

@router.get("/subjects", response_model=List[schemas.SubjectOut])
def get_subjects(
    department_id: Optional[int] = None,
    program_id: Optional[int] = None,
    academic_year_id: Optional[int] = None,
    semester_id: Optional[int] = None,
    search: Optional[str] = None,
    is_active: Optional[bool] = None,
    db: Session = Depends(get_db)
):
    q = db.query(models.Subject)
    if department_id:
        q = q.filter(models.Subject.department_id == department_id)
    if program_id:
        q = q.filter(models.Subject.program_id == program_id)
    if academic_year_id:
        q = q.filter(models.Subject.academic_year_id == academic_year_id)
    if semester_id:
        q = q.filter(models.Subject.semester_id == semester_id)
    if is_active is not None:
        q = q.filter(models.Subject.is_active == is_active)
    if search:
        s = f"%{search.strip()}%"
        q = q.filter(or_(models.Subject.name.ilike(s), models.Subject.code.ilike(s)))
    return q.order_by(models.Subject.name.asc()).all()


@router.get("/subjects/{subject_id}", response_model=schemas.SubjectOut)
def get_subject(subject_id: int, db: Session = Depends(get_db)):
    subj = db.query(models.Subject).filter(models.Subject.id == subject_id).first()
    if not subj:
        raise HTTPException(status_code=404, detail="Subject not found.")
    return subj


@router.post("/subjects", response_model=schemas.SubjectOut, status_code=status.HTTP_201_CREATED)
def create_subject(subj_in: schemas.SubjectCreate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    name = subj_in.name.strip()
    code = subj_in.code.strip()
    if not name or not code:
        raise HTTPException(status_code=400, detail="Subject name and code are required.")

    dept = db.query(models.Department).filter(models.Department.id == subj_in.department_id).first()
    if not dept:
        raise HTTPException(status_code=400, detail="Invalid department ID.")

    prog_name = None
    if subj_in.program_id:
        prog = db.query(models.Program).filter(models.Program.id == subj_in.program_id).first()
        if prog:
            prog_name = prog.name

    year_name = None
    if subj_in.academic_year_id:
        ay = db.query(models.AcademicYear).filter(models.AcademicYear.id == subj_in.academic_year_id).first()
        if ay:
            year_name = ay.name

    sem_name = None
    if subj_in.semester_id:
        sem = db.query(models.Semester).filter(models.Semester.id == subj_in.semester_id).first()
        if sem:
            sem_name = sem.name

    subj = models.Subject(
        name=name,
        code=code,
        department_id=dept.id,
        department_name=dept.name,
        program_id=subj_in.program_id,
        program_name=prog_name,
        academic_year_id=subj_in.academic_year_id,
        year_name=year_name,
        semester_id=subj_in.semester_id,
        semester_name=sem_name,
        is_active=subj_in.is_active,
        created_at=datetime.utcnow(),
        updated_at=datetime.utcnow()
    )
    db.add(subj)
    db.commit()
    db.refresh(subj)
    return subj


@router.put("/subjects/{subject_id}", response_model=schemas.SubjectOut)
def update_subject(subject_id: int, subj_in: schemas.SubjectUpdate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    subj = db.query(models.Subject).filter(models.Subject.id == subject_id).first()
    if not subj:
        raise HTTPException(status_code=404, detail="Subject not found.")

    if subj_in.name is not None:
        subj.name = subj_in.name.strip()
    if subj_in.code is not None:
        subj.code = subj_in.code.strip()
    if subj_in.department_id is not None:
        dept = db.query(models.Department).filter(models.Department.id == subj_in.department_id).first()
        if dept:
            subj.department_id = dept.id
            subj.department_name = dept.name
    if subj_in.program_id is not None:
        prog = db.query(models.Program).filter(models.Program.id == subj_in.program_id).first()
        subj.program_id = prog.id if prog else None
        subj.program_name = prog.name if prog else None
    if subj_in.academic_year_id is not None:
        ay = db.query(models.AcademicYear).filter(models.AcademicYear.id == subj_in.academic_year_id).first()
        subj.academic_year_id = ay.id if ay else None
        subj.year_name = ay.name if ay else None
    if subj_in.semester_id is not None:
        sem = db.query(models.Semester).filter(models.Semester.id == subj_in.semester_id).first()
        subj.semester_id = sem.id if sem else None
        subj.semester_name = sem.name if sem else None
    if subj_in.is_active is not None:
        subj.is_active = subj_in.is_active

    subj.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(subj)
    return subj


@router.patch("/subjects/{subject_id}/status", response_model=schemas.SubjectOut)
def toggle_subject_status(subject_id: int, status_in: schemas.QuestionStatusUpdate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    subj = db.query(models.Subject).filter(models.Subject.id == subject_id).first()
    if not subj:
        raise HTTPException(status_code=404, detail="Subject not found.")
    subj.is_active = status_in.is_active
    subj.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(subj)
    return subj


@router.delete("/subjects/{subject_id}")
def delete_subject(subject_id: int, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    subj = db.query(models.Subject).filter(models.Subject.id == subject_id).first()
    if not subj:
        raise HTTPException(status_code=404, detail="Subject not found.")

    # Safe delete check: verify if any questions or exams reference this subject
    q_count = db.query(models.Question).filter(models.Question.subject_id == subject_id).count()
    exam_count = db.query(models.Exam).filter(models.Exam.subject_id == subject_id).count()

    if q_count > 0 or exam_count > 0:
        raise HTTPException(
            status_code=400,
            detail=f"Cannot delete subject '{subj.name}' because it is referenced by {q_count} question(s) and {exam_count} exam(s). Please deactivate it instead to preserve academic records."
        )

    db.delete(subj)
    db.commit()
    return {"message": f"Subject '{subj.name}' deleted successfully", "id": subject_id}


# =========================================================================
# 4. STUDENT MANAGEMENT & BULK IMPORT
# =========================================================================

@router.get("/admin/students/list", response_model=List[schemas.StudentOut])
def list_students(
    department: Optional[str] = None,
    program: Optional[str] = None,
    year: Optional[str] = None,
    semester: Optional[str] = None,
    search: Optional[str] = None,
    is_active: Optional[bool] = None,
    admin: str = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    q = db.query(models.Student)
    if department:
        q = q.filter(models.Student.department == department)
    if program:
        q = q.filter(models.Student.program == program)
    if year:
        q = q.filter(models.Student.year == year)
    if semester:
        q = q.filter(models.Student.semester == semester)
    if is_active is not None:
        q = q.filter(models.Student.is_active == is_active)
    if search:
        s = f"%{search.strip()}%"
        q = q.filter(or_(models.Student.enrollment_no.ilike(s), models.Student.name.ilike(s)))
    return q.order_by(models.Student.enrollment_no.asc()).all()


@router.post("/admin/students", response_model=schemas.StudentOut, status_code=status.HTTP_201_CREATED)
def admin_create_student(student_in: schemas.StudentCreate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    enrollment = student_in.enrollment_no.strip()
    name = student_in.name.strip()
    dept = student_in.department.strip()

    if not enrollment or not name or not dept:
        raise HTTPException(status_code=400, detail="Enrollment No, Name, and Department are required.")

    existing = db.query(models.Student).filter(models.Student.enrollment_no == enrollment).first()
    if existing:
        raise HTTPException(status_code=400, detail=f"Student with Enrollment No '{enrollment}' already exists.")

    raw_password = student_in.password.strip() if student_in.password else enrollment
    hashed_pwd = hash_password(raw_password)

    student = models.Student(
        enrollment_no=enrollment,
        name=name,
        gender=student_in.gender.strip() if student_in.gender else None,
        department=dept,
        program=student_in.program.strip() if student_in.program else "UG",
        year=student_in.year.strip() if student_in.year else "1st Year",
        semester=student_in.semester.strip() if student_in.semester else "Semester 1",
        password_hash=hashed_pwd,
        is_active=True,
        created_at=datetime.utcnow()
    )
    db.add(student)
    db.commit()
    db.refresh(student)
    return student


@router.put("/admin/students/{student_id}", response_model=schemas.StudentOut)
def admin_update_student(student_id: int, student_in: schemas.StudentUpdate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    student = db.query(models.Student).filter(models.Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found.")

    if student_in.name is not None:
        student.name = student_in.name.strip()
    if student_in.department is not None:
        student.department = student_in.department.strip()
    if student_in.gender is not None:
        student.gender = student_in.gender.strip()
    if student_in.program is not None:
        student.program = student_in.program.strip()
    if student_in.year is not None:
        student.year = student_in.year.strip()
    if student_in.semester is not None:
        student.semester = student_in.semester.strip()
    if student_in.password is not None and student_in.password.strip():
        student.password_hash = hash_password(student_in.password.strip())
    if student_in.is_active is not None:
        student.is_active = student_in.is_active

    student.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(student)
    return student


@router.patch("/admin/students/{student_id}/status", response_model=schemas.StudentOut)
def toggle_student_status(student_id: int, status_in: schemas.QuestionStatusUpdate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    student = db.query(models.Student).filter(models.Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found.")
    student.is_active = status_in.is_active
    student.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(student)
    return student


@router.post("/admin/students/bulk-import")
async def bulk_import_students(
    file: UploadFile = File(...),
    department: str = Form(...),
    program: str = Form("UG"),
    year: str = Form("1st Year"),
    semester: str = Form("Semester 1"),
    admin: str = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    if not department.strip():
        raise HTTPException(status_code=400, detail="Department is required for student import.")

    contents = await file.read()
    filename = file.filename.lower() if file.filename else ""
    parsed_records = []

    try:
        if filename.endswith(".csv") or file.content_type == "text/csv":
            text = contents.decode("utf-8-sig", errors="ignore")
            reader = csv.reader(io.StringIO(text))
            rows = list(reader)
            if rows:
                header = [h.strip().lower() for h in rows[0]]
                roll_idx, name_idx, gender_idx = 0, 1, 2
                start_row = 0
                # Detect header
                for i, col in enumerate(header):
                    if any(k in col for k in ["roll", "enroll"]):
                        roll_idx = i
                        start_row = 1
                    elif any(k in col for k in ["name", "student"]):
                        name_idx = i
                        start_row = 1
                    elif "gender" in col or "sex" in col:
                        gender_idx = i
                        start_row = 1

                for row in rows[start_row:]:
                    if not row or len(row) <= roll_idx:
                        continue
                    r_no = row[roll_idx].strip()
                    nm = row[name_idx].strip() if len(row) > name_idx else ""
                    gnd = row[gender_idx].strip() if len(row) > gender_idx else None
                    if r_no and nm:
                        parsed_records.append((r_no, nm, gnd))

        elif filename.endswith((".xlsx", ".xls")):
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if rows:
                header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
                roll_idx, name_idx, gender_idx = 0, 1, 2
                start_row = 0
                for i, col in enumerate(header):
                    if any(k in col for k in ["roll", "enroll"]):
                        roll_idx = i
                        start_row = 1
                    elif any(k in col for k in ["name", "student"]):
                        name_idx = i
                        start_row = 1
                    elif "gender" in col or "sex" in col:
                        gender_idx = i
                        start_row = 1

                for row in rows[start_row:]:
                    if not row or len(row) <= roll_idx or row[roll_idx] is None:
                        continue
                    r_no = str(row[roll_idx]).strip()
                    nm = str(row[name_idx]).strip() if len(row) > name_idx and row[name_idx] is not None else ""
                    gnd = str(row[gender_idx]).strip() if len(row) > gender_idx and row[gender_idx] is not None else None
                    if r_no and nm:
                        parsed_records.append((r_no, nm, gnd))

        elif filename.endswith(".pdf") or file.content_type == "application/pdf":
            reader = PdfReader(io.BytesIO(contents))
            pdf_text = ""
            for page in reader.pages:
                txt = page.extract_text()
                if txt:
                    pdf_text += txt + "\n"

            # Parse lines for patterns like: [SrNo] RollNo Student Name [MALE|FEMALE]
            # Pattern: (BT26F05F001 or similar roll number) followed by name and gender
            lines = pdf_text.splitlines()
            for line in lines:
                line_str = line.strip()
                if not line_str:
                    continue
                # Regex match for roll number and name
                # E.g.: "1  BT26F05F001  ALAMWAR MOHD ZAID ABDUL WAHID  MALE"
                # or "BT26F05F001 JOHN DOE M"
                match = re.search(
                    r'(?:^|\s+)([A-Z0-9_-]{6,20})\s+([A-Za-z\s\.\'-]{3,50})(?:\s+(MALE|FEMALE|M|F|OTHER))?',
                    line_str,
                    re.IGNORECASE
                )
                if match:
                    r_no = match.group(1).strip()
                    # Skip common header strings if matched
                    if any(h in r_no.lower() for h in ["rollno", "enroll", "student", "number", "srno"]):
                        continue
                    nm = match.group(2).strip()
                    gnd = match.group(3).strip() if match.group(3) else None
                    if r_no and nm:
                        parsed_records.append((r_no, nm, gnd))
        else:
            raise HTTPException(status_code=400, detail="Unsupported file format. Please upload CSV, XLSX, or PDF.")

    except Exception as parse_err:
        raise HTTPException(status_code=400, detail=f"Failed to parse file: {str(parse_err)}")

    if not parsed_records:
        raise HTTPException(status_code=400, detail="No valid student records found in file. Ensure Roll No and Name columns are present.")

    imported_count = 0
    skipped_count = 0
    updated_count = 0

    default_hashed_pwd = hash_password("student123")
    now = datetime.utcnow()

    for r_no, nm, gnd in parsed_records:
        existing = db.query(models.Student).filter(models.Student.enrollment_no == r_no).first()
        if existing:
            # Update attributes safely
            existing.name = nm
            if gnd:
                existing.gender = gnd
            existing.department = department.strip()
            existing.program = program.strip()
            existing.year = year.strip()
            existing.semester = semester.strip()
            existing.is_active = True
            existing.updated_at = now
            updated_count += 1
        else:
            new_student = models.Student(
                enrollment_no=r_no,
                name=nm,
                gender=gnd,
                department=department.strip(),
                program=program.strip(),
                year=year.strip(),
                semester=semester.strip(),
                password_hash=default_hashed_pwd,
                is_active=True,
                created_at=now,
                updated_at=now
            )
            db.add(new_student)
            imported_count += 1

    db.commit()

    return {
        "message": f"Student import completed: {imported_count} new imported, {updated_count} existing updated.",
        "imported_count": imported_count,
        "updated_count": updated_count,
        "total_processed": len(parsed_records)
    }


# =========================================================================
# 5. SECURE STUDENT LOGIN & STUDENT APIS
# =========================================================================

@router.post("/student/login", response_model=schemas.StudentToken)
def student_login(credentials: schemas.StudentLogin, db: Session = Depends(get_db)):
    enrollment = credentials.enrollment_no.strip()
    plain_pwd = credentials.password.strip()

    if not enrollment or not plain_pwd:
        raise HTTPException(status_code=400, detail="Enrollment number and password are required.")

    student = db.query(models.Student).filter(models.Student.enrollment_no == enrollment).first()
    if not student:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid Enrollment Number or Password.")

    if not student.is_active:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Student account has been deactivated. Please contact admin.")

    # Validate password hash
    if not student.password_hash:
        # If student password was not set yet, fallback allow enrollment or default 'student123'
        if plain_pwd in [student.enrollment_no, "student123"]:
            student.password_hash = hash_password(plain_pwd)
            db.commit()
        else:
            raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid Enrollment Number or Password.")
    else:
        if not verify_password(plain_pwd, student.password_hash):
            raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid Enrollment Number or Password.")

    access_token = create_access_token(data={"sub": str(student.id), "enrollment_no": student.enrollment_no, "role": "student"})
    return {
        "access_token": access_token,
        "token_type": "bearer",
        "student": student
    }


@router.get("/student/profile", response_model=schemas.StudentOut)
def get_student_profile(current_student: models.Student = Depends(get_current_student)):
    return current_student


@router.get("/student/exams", response_model=List[schemas.StudentAvailableExam])
def get_student_available_exams(
    current_student: models.Student = Depends(get_current_student),
    db: Session = Depends(get_db)
):
    """
    CRITICAL REQUIREMENT: STUDENT-SPECIFIC EXAM ACCESS
    Only return exams that are assigned to this specific student via exam_students!
    """
    assigned_records = db.query(models.ExamStudent).filter(
        models.ExamStudent.student_id == current_student.id
    ).all()

    assigned_exam_ids = [ar.exam_id for ar in assigned_records]
    if not assigned_exam_ids:
        return []

    exams = db.query(models.Exam).filter(
        models.Exam.id.in_(assigned_exam_ids),
        models.Exam.is_active == True,
        models.Exam.status == "published"
    ).order_by(models.Exam.created_at.desc()).all()

    results = []
    now = datetime.utcnow()

    for ex in exams:
        # Check if student already has an attempt
        attempt = db.query(models.Attempt).filter(
            models.Attempt.exam_id == ex.id,
            models.Attempt.student_id == current_student.id
        ).order_by(models.Attempt.started_at.desc()).first()

        status_str = "available"
        attempt_id = None
        score = None
        percentage = None

        if attempt:
            attempt_id = attempt.id
            if attempt.status in ["completed", "timed_out"]:
                status_str = attempt.status
                score = attempt.score
                percentage = attempt.percentage
            elif attempt.status == "in_progress":
                if attempt.deadline_at and now >= attempt.deadline_at:
                    # Automatically finalize expired attempt
                    attempt.status = "timed_out"
                    attempt.completed_at = attempt.deadline_at
                    calculate_attempt_score(attempt, db)
                    db.commit()
                    status_str = "timed_out"
                    score = attempt.score
                    percentage = attempt.percentage
                else:
                    status_str = "in_progress"

        results.append({
            "id": ex.id,
            "title": ex.title,
            "code": ex.code,
            "subject_name": ex.subject_name or (ex.subject.name if ex.subject else "General"),
            "department_name": ex.department_name,
            "duration_minutes": ex.duration_minutes,
            "total_questions": ex.total_questions,
            "total_marks": ex.total_marks,
            "passing_percentage": ex.passing_percentage,
            "start_date": ex.start_date,
            "start_time": ex.start_time,
            "end_date": ex.end_date,
            "end_time": ex.end_time,
            "status": status_str,
            "attempt_id": attempt_id,
            "score": score,
            "percentage": percentage
        })

    return results


# =========================================================================
# 6. EXAM MANAGEMENT (ADMIN)
# =========================================================================

@router.get("/admin/exams", response_model=List[schemas.ExamOut])
def list_exams(
    department_id: Optional[int] = None,
    program_id: Optional[int] = None,
    academic_year_id: Optional[int] = None,
    semester_id: Optional[int] = None,
    subject_id: Optional[int] = None,
    status: Optional[str] = None,
    search: Optional[str] = None,
    admin: str = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    q = db.query(models.Exam)
    if department_id:
        q = q.filter(models.Exam.department_id == department_id)
    if program_id:
        q = q.filter(models.Exam.program_id == program_id)
    if academic_year_id:
        q = q.filter(models.Exam.academic_year_id == academic_year_id)
    if semester_id:
        q = q.filter(models.Exam.semester_id == semester_id)
    if subject_id:
        q = q.filter(models.Exam.subject_id == subject_id)
    if status:
        q = q.filter(models.Exam.status == status)
    if search:
        s = f"%{search.strip()}%"
        q = q.filter(or_(models.Exam.title.ilike(s), models.Exam.code.ilike(s)))

    exams = q.order_by(models.Exam.created_at.desc()).all()
    out = []
    for ex in exams:
        assigned_cnt = db.query(models.ExamStudent).filter(models.ExamStudent.exam_id == ex.id).count()
        attempts_cnt = db.query(models.Attempt).filter(models.Attempt.exam_id == ex.id).count()
        e_dict = schemas.ExamOut.model_validate(ex)
        e_dict.assigned_students_count = assigned_cnt
        e_dict.attempts_count = attempts_cnt
        out.append(e_dict)
    return out


@router.get("/admin/exams/{exam_id}", response_model=schemas.ExamOut)
def get_exam(exam_id: int, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    ex = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not ex:
        raise HTTPException(status_code=404, detail="Exam not found.")
    assigned_cnt = db.query(models.ExamStudent).filter(models.ExamStudent.exam_id == ex.id).count()
    attempts_cnt = db.query(models.Attempt).filter(models.Attempt.exam_id == ex.id).count()
    e_dict = schemas.ExamOut.model_validate(ex)
    e_dict.assigned_students_count = assigned_cnt
    e_dict.attempts_count = attempts_cnt
    return e_dict


@router.post("/admin/exams", response_model=schemas.ExamOut, status_code=status.HTTP_201_CREATED)
def create_exam(exam_in: schemas.ExamCreate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    title = exam_in.title.strip()
    if not title:
        raise HTTPException(status_code=400, detail="Exam title is required.")

    dept_name = None
    if exam_in.department_id:
        dept = db.query(models.Department).filter(models.Department.id == exam_in.department_id).first()
        dept_name = dept.name if dept else None

    subj_name = None
    if exam_in.subject_id:
        subj = db.query(models.Subject).filter(models.Subject.id == exam_in.subject_id).first()
        subj_name = subj.name if subj else None

    prog_name = None
    if exam_in.program_id:
        prog = db.query(models.Program).filter(models.Program.id == exam_in.program_id).first()
        prog_name = prog.name if prog else None

    year_name = None
    if exam_in.academic_year_id:
        ay = db.query(models.AcademicYear).filter(models.AcademicYear.id == exam_in.academic_year_id).first()
        year_name = ay.name if ay else None

    sem_name = None
    if exam_in.semester_id:
        sem = db.query(models.Semester).filter(models.Semester.id == exam_in.semester_id).first()
        sem_name = sem.name if sem else None

    total_marks = exam_in.total_marks or (exam_in.total_questions * exam_in.marks_per_question)

    exam = models.Exam(
        title=title,
        code=exam_in.code.strip() if exam_in.code else None,
        department_id=exam_in.department_id,
        department_name=dept_name,
        program_id=exam_in.program_id,
        program_name=prog_name,
        academic_year_id=exam_in.academic_year_id,
        year_name=year_name,
        semester_id=exam_in.semester_id,
        semester_name=sem_name,
        subject_id=exam_in.subject_id,
        subject_name=subj_name,
        start_date=exam_in.start_date,
        start_time=exam_in.start_time,
        end_date=exam_in.end_date,
        end_time=exam_in.end_time,
        duration_minutes=exam_in.duration_minutes,
        total_questions=exam_in.total_questions,
        marks_per_question=exam_in.marks_per_question,
        total_marks=total_marks,
        passing_percentage=exam_in.passing_percentage,
        selection_mode=exam_in.selection_mode,
        status=exam_in.status,
        is_active=True,
        created_at=datetime.utcnow(),
        updated_at=datetime.utcnow()
    )
    db.add(exam)
    db.commit()
    db.refresh(exam)

    # If manual questions selection
    if exam_in.selection_mode == "manual" and exam_in.selected_question_ids:
        for idx, qid in enumerate(exam_in.selected_question_ids, start=1):
            eq = models.ExamQuestion(exam_id=exam.id, question_id=qid, order_index=idx)
            db.add(eq)
        db.commit()

    e_dict = schemas.ExamOut.model_validate(exam)
    e_dict.assigned_students_count = 0
    e_dict.attempts_count = 0
    return e_dict


@router.put("/admin/exams/{exam_id}", response_model=schemas.ExamOut)
def update_exam(exam_id: int, exam_in: schemas.ExamUpdate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found.")

    if exam_in.title is not None:
        exam.title = exam_in.title.strip()
    if exam_in.code is not None:
        exam.code = exam_in.code.strip()
    if exam_in.department_id is not None:
        dept = db.query(models.Department).filter(models.Department.id == exam_in.department_id).first()
        exam.department_id = dept.id if dept else None
        exam.department_name = dept.name if dept else None
    if exam_in.subject_id is not None:
        subj = db.query(models.Subject).filter(models.Subject.id == exam_in.subject_id).first()
        exam.subject_id = subj.id if subj else None
        exam.subject_name = subj.name if subj else None
    if exam_in.program_id is not None:
        prog = db.query(models.Program).filter(models.Program.id == exam_in.program_id).first()
        exam.program_id = prog.id if prog else None
        exam.program_name = prog.name if prog else None
    if exam_in.academic_year_id is not None:
        ay = db.query(models.AcademicYear).filter(models.AcademicYear.id == exam_in.academic_year_id).first()
        exam.academic_year_id = ay.id if ay else None
        exam.year_name = ay.name if ay else None
    if exam_in.semester_id is not None:
        sem = db.query(models.Semester).filter(models.Semester.id == exam_in.semester_id).first()
        exam.semester_id = sem.id if sem else None
        exam.semester_name = sem.name if sem else None
    if exam_in.start_date is not None:
        exam.start_date = exam_in.start_date
    if exam_in.start_time is not None:
        exam.start_time = exam_in.start_time
    if exam_in.end_date is not None:
        exam.end_date = exam_in.end_date
    if exam_in.end_time is not None:
        exam.end_time = exam_in.end_time
    if exam_in.duration_minutes is not None:
        exam.duration_minutes = exam_in.duration_minutes
    if exam_in.total_questions is not None:
        exam.total_questions = exam_in.total_questions
    if exam_in.marks_per_question is not None:
        exam.marks_per_question = exam_in.marks_per_question
    if exam_in.total_marks is not None:
        exam.total_marks = exam_in.total_marks
    if exam_in.passing_percentage is not None:
        exam.passing_percentage = exam_in.passing_percentage
    if exam_in.selection_mode is not None:
        exam.selection_mode = exam_in.selection_mode
    if exam_in.status is not None:
        exam.status = exam_in.status
    if exam_in.is_active is not None:
        exam.is_active = exam_in.is_active

    if exam_in.selected_question_ids is not None:
        db.query(models.ExamQuestion).filter(models.ExamQuestion.exam_id == exam.id).delete()
        for idx, qid in enumerate(exam_in.selected_question_ids, start=1):
            db.add(models.ExamQuestion(exam_id=exam.id, question_id=qid, order_index=idx))

    exam.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(exam)

    assigned_cnt = db.query(models.ExamStudent).filter(models.ExamStudent.exam_id == exam.id).count()
    attempts_cnt = db.query(models.Attempt).filter(models.Attempt.exam_id == exam.id).count()
    e_dict = schemas.ExamOut.model_validate(exam)
    e_dict.assigned_students_count = assigned_cnt
    e_dict.attempts_count = attempts_cnt
    return e_dict


@router.patch("/admin/exams/{exam_id}/status", response_model=schemas.ExamOut)
def set_exam_status(exam_id: int, status_in: schemas.QuestionStatusUpdate, status_name: Optional[str] = None, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found.")
    if status_name:
        exam.status = status_name
    exam.is_active = status_in.is_active
    exam.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(exam)
    assigned_cnt = db.query(models.ExamStudent).filter(models.ExamStudent.exam_id == exam.id).count()
    attempts_cnt = db.query(models.Attempt).filter(models.Attempt.exam_id == exam.id).count()
    e_dict = schemas.ExamOut.model_validate(exam)
    e_dict.assigned_students_count = assigned_cnt
    e_dict.attempts_count = attempts_cnt
    return e_dict


@router.get("/admin/exams/{exam_id}/students", response_model=List[schemas.AssignedStudentOut])
def get_exam_assigned_students(exam_id: int, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found.")

    assignments = db.query(models.ExamStudent).filter(models.ExamStudent.exam_id == exam_id).all()
    results = []
    for a in assignments:
        st = a.student
        if not st:
            continue
        attempt = db.query(models.Attempt).filter(
            models.Attempt.exam_id == exam_id,
            models.Attempt.student_id == st.id
        ).order_by(models.Attempt.started_at.desc()).first()

        results.append({
            "student_id": st.id,
            "enrollment_no": st.enrollment_no,
            "name": st.name,
            "department": st.department,
            "gender": st.gender,
            "program": st.program,
            "year": st.year,
            "semester": st.semester,
            "assigned_at": a.assigned_at,
            "attempt_id": attempt.id if attempt else None,
            "status": attempt.status if attempt else "unattempted",
            "score": attempt.score if attempt else None,
            "percentage": attempt.percentage if attempt else None
        })
    return results


@router.post("/admin/exams/{exam_id}/assign")
def assign_students_to_exam(
    exam_id: int,
    assign_in: schemas.AssignStudentsRequest,
    admin: str = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found.")

    now = datetime.utcnow()
    added_count = 0
    for sid in assign_in.student_ids:
        exists = db.query(models.ExamStudent).filter(
            models.ExamStudent.exam_id == exam_id,
            models.ExamStudent.student_id == sid
        ).first()
        if not exists:
            db.add(models.ExamStudent(exam_id=exam_id, student_id=sid, assigned_at=now))
            added_count += 1
    db.commit()

    total_assigned = db.query(models.ExamStudent).filter(models.ExamStudent.exam_id == exam_id).count()
    return {"message": f"Successfully assigned {added_count} students to '{exam.title}'.", "total_assigned": total_assigned}


@router.delete("/admin/exams/{exam_id}/students/{student_id}")
def unassign_student_from_exam(exam_id: int, student_id: int, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    rec = db.query(models.ExamStudent).filter(
        models.ExamStudent.exam_id == exam_id,
        models.ExamStudent.student_id == student_id
    ).first()
    if not rec:
        raise HTTPException(status_code=404, detail="Student is not assigned to this exam.")
    db.delete(rec)
    db.commit()
    return {"message": "Student unassigned from exam successfully."}


@router.delete("/admin/exams/{exam_id}")
def delete_exam(exam_id: int, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found.")

    attempt_cnt = db.query(models.Attempt).filter(models.Attempt.exam_id == exam_id).count()
    if attempt_cnt > 0:
        raise HTTPException(
            status_code=400,
            detail=f"Cannot delete exam '{exam.title}' because it has {attempt_cnt} recorded attempt(s). Please close or deactivate the exam instead."
        )

    db.delete(exam)
    db.commit()
    return {"message": f"Exam '{exam.title}' deleted successfully."}


# =========================================================================
# 7. CRITICAL STUDENT EXAM START & BACKEND AUTHORIZATION
# =========================================================================

@router.post("/student/exams/{exam_id}/start", response_model=schemas.StartAttemptResponse)
def student_start_exam(
    exam_id: int,
    current_student: models.Student = Depends(get_current_student),
    db: Session = Depends(get_db)
):
    """
    CRITICAL REQUIREMENT: STUDENT-SPECIFIC EXAM ACCESS
    Enforces authorization on the backend.
    Unassigned student receives 403 Forbidden.
    """
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found.")

    if not exam.is_active:
        raise HTTPException(status_code=400, detail="This exam is currently inactive.")

    if exam.status != "published":
        raise HTTPException(status_code=400, detail=f"Exam is currently {exam.status} and not open for attempts.")

    # BACKEND AUTHORIZATION CHECK: Is student assigned?
    assignment = db.query(models.ExamStudent).filter(
        models.ExamStudent.exam_id == exam_id,
        models.ExamStudent.student_id == current_student.id
    ).first()

    if not assignment:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Exam is not assigned to this student."
        )

    now = datetime.utcnow()

    # Check date/time window if defined
    if exam.start_date:
        today_str = now.strftime("%Y-%m-%d")
        if today_str < exam.start_date:
            raise HTTPException(status_code=400, detail=f"Exam has not started yet. Starts on {exam.start_date}.")
        if exam.end_date and today_str > exam.end_date:
            raise HTTPException(status_code=400, detail=f"Exam availability has ended on {exam.end_date}.")

    # Check for existing attempts
    existing_attempt = db.query(models.Attempt).filter(
        models.Attempt.exam_id == exam_id,
        models.Attempt.student_id == current_student.id
    ).order_by(models.Attempt.started_at.desc()).first()

    if existing_attempt:
        if existing_attempt.status in ["completed", "timed_out"]:
            raise HTTPException(status_code=400, detail="You have already completed this exam.")
        
        # In progress: check deadline
        if existing_attempt.deadline_at and now >= existing_attempt.deadline_at:
            existing_attempt.status = "timed_out"
            existing_attempt.completed_at = existing_attempt.deadline_at
            calculate_attempt_score(existing_attempt, db)
            db.commit()
            raise HTTPException(status_code=400, detail="Your exam attempt time limit has expired.")

        # Resume existing attempt
        return {
            "attempt_id": existing_attempt.id,
            "student": current_student,
            "message": "Resumed ongoing exam attempt.",
            "exam_title": exam.title,
            "deadline_at": existing_attempt.deadline_at,
            "time_limit_minutes": existing_attempt.time_limit_minutes,
            "total_questions": existing_attempt.total_questions
        }

    # Create new attempt with server-side timer
    started_at = now
    deadline_at = started_at + timedelta(minutes=exam.duration_minutes)

    attempt = models.Attempt(
        student_id=current_student.id,
        exam_id=exam.id,
        exam_title_snapshot=exam.title,
        subject_name_snapshot=exam.subject_name or (exam.subject.name if exam.subject else "General"),
        total_questions=exam.total_questions,
        question_count_snapshot=exam.total_questions,
        time_limit_minutes=exam.duration_minutes,
        started_at=started_at,
        deadline_at=deadline_at,
        status="in_progress",
        tab_switch_count=0,
        fullscreen_exit_count=0,
        copy_count=0,
        paste_count=0
    )
    db.add(attempt)
    db.commit()
    db.refresh(attempt)

    # Question Selection:
    # 1. Manual mode: use exam_questions
    chosen_questions = []
    if exam.selection_mode == "manual":
        eqs = db.query(models.ExamQuestion).filter(models.ExamQuestion.exam_id == exam.id).order_by(models.ExamQuestion.order_index.asc()).all()
        for eq in eqs:
            if eq.question and eq.question.is_active:
                chosen_questions.append(eq.question)

    # 2. Random mode or fallback: pick from subject's question bank matching department/subject
    if len(chosen_questions) < exam.total_questions:
        q_filter = db.query(models.Question).filter(models.Question.is_active == True)
        if exam.subject_id:
            q_filter = q_filter.filter(models.Question.subject_id == exam.subject_id)
        elif exam.department_id:
            q_filter = q_filter.filter(models.Question.department_id == exam.department_id)

        available_qs = q_filter.all()
        # If not enough subject-specific questions, fallback to any active questions so exam can proceed
        if len(available_qs) < exam.total_questions:
            fallback_qs = db.query(models.Question).filter(models.Question.is_active == True).all()
            for fq in fallback_qs:
                if fq not in available_qs:
                    available_qs.append(fq)

        needed = exam.total_questions - len(chosen_questions)
        already_ids = {q.id for q in chosen_questions}
        pool = [q for q in available_qs if q.id not in already_ids]
        sampled = random.sample(pool, min(needed, len(pool)))
        chosen_questions.extend(sampled)

    # Create student answers snapshot
    student_answers = []
    for idx, q in enumerate(chosen_questions[:exam.total_questions], start=1):
        sa = models.StudentAnswer(
            attempt_id=attempt.id,
            question_number=idx,
            question_id=q.id,
            question_text=q.question_text,
            option_a=q.option_a,
            option_b=q.option_b,
            option_c=q.option_c,
            option_d=q.option_d,
            selected_answer=None,
            correct_answer=q.correct_answer,
            marks=q.marks or exam.marks_per_question or 1,
            is_correct=False,
            answered_at=started_at
        )
        student_answers.append(sa)

    db.bulk_save_objects(student_answers)
    db.commit()

    return {
        "attempt_id": attempt.id,
        "student": current_student,
        "message": "Exam attempt started successfully.",
        "exam_title": exam.title,
        "deadline_at": deadline_at,
        "time_limit_minutes": exam.duration_minutes,
        "total_questions": len(student_answers)
    }


# =========================================================================
# 8. SECURITY & CHEATING AUDIT LOGGING
# =========================================================================

@router.post("/attempts/{attempt_id}/security-log", response_model=schemas.SecurityLogOut)
def record_security_log(
    attempt_id: int,
    log_in: schemas.SecurityLogCreate,
    db: Session = Depends(get_db)
):
    attempt = db.query(models.Attempt).filter(models.Attempt.id == attempt_id).first()
    if not attempt:
        raise HTTPException(status_code=404, detail="Attempt not found.")

    now = datetime.utcnow()

    # Update summary counters on attempt
    etype = log_in.event_type.lower()
    if "tab" in etype:
        attempt.tab_switch_count = (attempt.tab_switch_count or 0) + 1
    elif "fullscreen" in etype:
        attempt.fullscreen_exit_count = (attempt.fullscreen_exit_count or 0) + 1
    elif "copy" in etype:
        attempt.copy_count = (attempt.copy_count or 0) + 1
    elif "paste" in etype:
        attempt.paste_count = (attempt.paste_count or 0) + 1

    sec_log = models.ExamSecurityLog(
        attempt_id=attempt.id,
        student_id=attempt.student_id,
        event_type=log_in.event_type,
        details=log_in.details,
        occurred_at=now
    )
    db.add(sec_log)
    db.commit()
    db.refresh(sec_log)
    return sec_log


# =========================================================================
# 9. GENERAL & EXISTING ATTEMPT FLOW (100% BACKWARD COMPATIBLE)
# =========================================================================

@router.get("/questions", response_model=List[schemas.QuestionOut])
def get_questions(attempt_id: Optional[int] = Query(None), db: Session = Depends(get_db)):
    if attempt_id:
        attempt = db.query(models.Attempt).filter(models.Attempt.id == attempt_id).first()
        if not attempt:
            raise HTTPException(status_code=404, detail="Attempt not found.")
        
        student_answers = db.query(models.StudentAnswer).filter(
            models.StudentAnswer.attempt_id == attempt_id
        ).order_by(models.StudentAnswer.question_number.asc()).all()

        sanitized = []
        for sa in student_answers:
            opts = {
                "A": sa.option_a or "",
                "B": sa.option_b or "",
                "C": sa.option_c or "",
                "D": sa.option_d or ""
            }
            if not any(opts.values()) and sa.question_number in QUESTIONS_BY_ID:
                opts = QUESTIONS_BY_ID[sa.question_number].get("options", {})

            q_text = sa.question_text
            if not q_text and sa.question_number in QUESTIONS_BY_ID:
                q_text = QUESTIONS_BY_ID[sa.question_number].get("question", "")

            sanitized.append({
                "id": sa.question_number,
                "question": q_text or f"Question {sa.question_number}",
                "options": opts
            })
        return sanitized

    asst_setting = get_or_create_settings(db)
    target_count = asst_setting.question_count

    active_questions = db.query(models.Question).filter(
        models.Question.is_active == True
    ).order_by(models.Question.id.asc()).limit(target_count).all()

    if active_questions:
        sanitized = []
        for q in active_questions:
            sanitized.append({
                "id": q.id,
                "question": q.question_text,
                "options": {
                    "A": q.option_a,
                    "B": q.option_b,
                    "C": q.option_c,
                    "D": q.option_d
                }
            })
        return sanitized

    sanitized = []
    for q in FIXED_QUESTIONS[:target_count]:
        sanitized.append({
            "id": q["id"],
            "question": q["question"],
            "options": q["options"]
        })
    return sanitized


@router.post("/students", response_model=schemas.StudentOut)
def create_student(student_data: schemas.StudentCreate, db: Session = Depends(get_db)):
    enrollment = student_data.enrollment_no.strip()
    name = student_data.name.strip()
    dept = student_data.department.strip()

    if not enrollment or not name or not dept:
        raise HTTPException(status_code=400, detail="All fields (Enrollment No, Name, Department) are required.")

    student = db.query(models.Student).filter(models.Student.enrollment_no == enrollment).first()
    default_pwd_hash = hash_password(student_data.password.strip() if student_data.password else enrollment)

    if not student:
        student = models.Student(
            enrollment_no=enrollment,
            name=name,
            department=dept,
            gender=student_data.gender,
            program=student_data.program or "UG",
            year=student_data.year or "1st Year",
            semester=student_data.semester or "Semester 1",
            password_hash=default_pwd_hash,
            is_active=True
        )
        db.add(student)
        db.commit()
        db.refresh(student)
    else:
        student.name = name
        student.department = dept
        if student_data.gender:
            student.gender = student_data.gender
        if student_data.program:
            student.program = student_data.program
        if student_data.year:
            student.year = student_data.year
        if student_data.semester:
            student.semester = student_data.semester
        if not student.password_hash:
            student.password_hash = default_pwd_hash
        db.commit()
        db.refresh(student)

    return student


@router.post("/attempts/start", response_model=schemas.StartAttemptResponse)
def start_attempt(student_data: schemas.StudentCreate, db: Session = Depends(get_db)):
    """Existing DSA MCQ assessment attempt start flow."""
    enrollment = student_data.enrollment_no.strip()
    name = student_data.name.strip()
    dept = student_data.department.strip()

    if not enrollment or not name or not dept:
        raise HTTPException(status_code=400, detail="All fields (Enrollment No, Name, Department) are required.")

    student = db.query(models.Student).filter(models.Student.enrollment_no == enrollment).first()
    default_pwd_hash = hash_password(enrollment)
    if not student:
        student = models.Student(
            enrollment_no=enrollment,
            name=name,
            department=dept,
            password_hash=default_pwd_hash,
            is_active=True
        )
        db.add(student)
        db.commit()
        db.refresh(student)
    else:
        student.name = name
        student.department = dept
        if not student.password_hash:
            student.password_hash = default_pwd_hash
        db.commit()
        db.refresh(student)

    # Check for in-progress attempt
    existing_attempt = db.query(models.Attempt).filter(
        models.Attempt.student_id == student.id,
        models.Attempt.status == "in_progress",
        models.Attempt.exam_id == None
    ).first()

    now = datetime.utcnow()
    if existing_attempt:
        if existing_attempt.deadline_at and now >= existing_attempt.deadline_at:
            existing_attempt.status = "timed_out"
            existing_attempt.completed_at = existing_attempt.deadline_at
            calculate_attempt_score(existing_attempt, db)
            db.commit()
        else:
            return {
                "attempt_id": existing_attempt.id,
                "student": student,
                "message": "Continuing your test in progress.",
                "deadline_at": existing_attempt.deadline_at,
                "time_limit_minutes": existing_attempt.time_limit_minutes or 60,
                "total_questions": existing_attempt.question_count_snapshot or 50
            }

    # Snapshot current assessment settings
    setting = get_or_create_settings(db)
    target_q_count = setting.question_count
    target_time_limit = setting.time_limit_minutes

    active_questions = db.query(models.Question).filter(
        models.Question.is_active == True
    ).order_by(models.Question.id.asc()).limit(target_q_count).all()

    actual_q_count = len(active_questions)
    if actual_q_count == 0 and FIXED_QUESTIONS:
        actual_q_count = min(target_q_count, len(FIXED_QUESTIONS))

    deadline_at = now + timedelta(minutes=target_time_limit)

    attempt = models.Attempt(
        student_id=student.id,
        exam_id=None,
        exam_title_snapshot="DSA 50-Question Assessment",
        subject_name_snapshot="Data Structures & Algorithms",
        total_questions=actual_q_count,
        question_count_snapshot=actual_q_count,
        time_limit_minutes=target_time_limit,
        score=0,
        percentage=0.0,
        correct_answers=0,
        wrong_answers=0,
        started_at=now,
        deadline_at=deadline_at,
        status="in_progress",
        tab_switch_count=0,
        fullscreen_exit_count=0
    )
    db.add(attempt)
    db.commit()
    db.refresh(attempt)

    # Create student answers snapshot
    student_answers = []
    if active_questions:
        for idx, q in enumerate(active_questions, start=1):
            sa = models.StudentAnswer(
                attempt_id=attempt.id,
                question_number=idx,
                question_id=q.id,
                question_text=q.question_text,
                option_a=q.option_a,
                option_b=q.option_b,
                option_c=q.option_c,
                option_d=q.option_d,
                selected_answer=None,
                correct_answer=q.correct_answer,
                marks=q.marks or 1,
                is_correct=False,
                answered_at=now
            )
            student_answers.append(sa)
    else:
        for q in FIXED_QUESTIONS[:actual_q_count]:
            opts = q.get("options", {})
            sa = models.StudentAnswer(
                attempt_id=attempt.id,
                question_number=q["id"],
                question_id=None,
                question_text=q["question"],
                option_a=opts.get("A", ""),
                option_b=opts.get("B", ""),
                option_c=opts.get("C", ""),
                option_d=opts.get("D", ""),
                selected_answer=None,
                correct_answer=q["correct_answer"],
                marks=1,
                is_correct=False,
                answered_at=now
            )
            student_answers.append(sa)

    db.bulk_save_objects(student_answers)
    db.commit()

    return {
        "attempt_id": attempt.id,
        "student": student,
        "message": "Assessment attempt started.",
        "deadline_at": deadline_at,
        "time_limit_minutes": target_time_limit,
        "total_questions": actual_q_count
    }


@router.get("/attempts/{attempt_id}/status", response_model=schemas.AttemptStatusResponse)
def get_attempt_status(attempt_id: int, db: Session = Depends(get_db)):
    attempt = db.query(models.Attempt).filter(models.Attempt.id == attempt_id).first()
    if not attempt:
        raise HTTPException(status_code=404, detail="Attempt not found.")

    now = datetime.utcnow()

    # Server-side timer & auto-submit check
    if attempt.status == "in_progress" and attempt.deadline_at and now >= attempt.deadline_at:
        attempt.status = "timed_out"
        attempt.completed_at = attempt.deadline_at
        calculate_attempt_score(attempt, db)
        db.commit()
        db.refresh(attempt)

    remaining_seconds = 0
    if attempt.status == "in_progress" and attempt.deadline_at:
        remaining_seconds = max(0, int((attempt.deadline_at - now).total_seconds()))

    student_answers = db.query(models.StudentAnswer).filter(
        models.StudentAnswer.attempt_id == attempt_id
    ).all()

    answers_map = {sa.question_number: sa.selected_answer for sa in student_answers}

    return {
        "attempt_id": attempt.id,
        "status": attempt.status,
        "exam_id": attempt.exam_id,
        "exam_title": attempt.exam_title_snapshot,
        "subject_name": attempt.subject_name_snapshot,
        "started_at": attempt.started_at,
        "deadline_at": attempt.deadline_at,
        "completed_at": attempt.completed_at,
        "time_limit_minutes": attempt.time_limit_minutes or 60,
        "total_questions": attempt.question_count_snapshot or attempt.total_questions or 50,
        "remaining_seconds": remaining_seconds,
        "current_server_time": now,
        "tab_switch_count": attempt.tab_switch_count or 0,
        "fullscreen_exit_count": attempt.fullscreen_exit_count or 0,
        "answers": answers_map
    }


@router.post("/attempts/{attempt_id}/answers")
def save_student_answer(
    attempt_id: int,
    payload: Dict,
    db: Session = Depends(get_db)
):
    attempt = db.query(models.Attempt).filter(models.Attempt.id == attempt_id).first()
    if not attempt:
        raise HTTPException(status_code=404, detail="Attempt not found.")

    now = datetime.utcnow()
    # Check if timer expired
    if attempt.status != "in_progress" or (attempt.deadline_at and now >= attempt.deadline_at):
        if attempt.status == "in_progress":
            attempt.status = "timed_out"
            attempt.completed_at = attempt.deadline_at
            calculate_attempt_score(attempt, db)
            db.commit()
        raise HTTPException(status_code=400, detail="Cannot save answers: Attempt has been completed or timed out.")

    answers_list = payload.get("answers", [])
    if not answers_list and "question_number" in payload:
        answers_list = [payload]

    updated_count = 0
    for ans in answers_list:
        q_num = ans.get("question_number")
        sel = ans.get("selected_answer")
        if sel is not None:
            sel = sel.strip().upper()
            if sel not in ["A", "B", "C", "D", ""]:
                sel = None

        sa = db.query(models.StudentAnswer).filter(
            models.StudentAnswer.attempt_id == attempt_id,
            models.StudentAnswer.question_number == q_num
        ).first()

        if sa:
            sa.selected_answer = sel
            sa.is_correct = (sel == sa.correct_answer)
            sa.answered_at = now
            updated_count += 1

    db.commit()
    return {"status": "success", "updated_count": updated_count}


@router.post("/attempts/{attempt_id}/submit", response_model=schemas.TestResultOut)
def submit_attempt(attempt_id: int, db: Session = Depends(get_db)):
    attempt = db.query(models.Attempt).filter(models.Attempt.id == attempt_id).first()
    if not attempt:
        raise HTTPException(status_code=404, detail="Attempt not found.")

    now = datetime.utcnow()
    if attempt.status == "in_progress":
        is_timed_out = attempt.deadline_at and now >= attempt.deadline_at
        attempt.status = "timed_out" if is_timed_out else "completed"
        attempt.submitted_at = now
        attempt.completed_at = now
        calculate_attempt_score(attempt, db)
        db.commit()
        db.refresh(attempt)

    return get_attempt_result(attempt_id, db)


@router.get("/attempts/{attempt_id}/result", response_model=schemas.TestResultOut)
def get_attempt_result(attempt_id: int, db: Session = Depends(get_db)):
    attempt = db.query(models.Attempt).filter(models.Attempt.id == attempt_id).first()
    if not attempt:
        raise HTTPException(status_code=404, detail="Attempt not found.")

    student = db.query(models.Student).filter(models.Student.id == attempt.student_id).first()
    student_answers = db.query(models.StudentAnswer).filter(
        models.StudentAnswer.attempt_id == attempt_id
    ).order_by(models.StudentAnswer.question_number.asc()).all()

    security_logs = db.query(models.ExamSecurityLog).filter(
        models.ExamSecurityLog.attempt_id == attempt_id
    ).order_by(models.ExamSecurityLog.occurred_at.asc()).all()

    answers_out = []
    for sa in student_answers:
        opts = {
            "A": sa.option_a or "",
            "B": sa.option_b or "",
            "C": sa.option_c or "",
            "D": sa.option_d or ""
        }
        if not any(opts.values()) and sa.question_number in QUESTIONS_BY_ID:
            opts = QUESTIONS_BY_ID[sa.question_number].get("options", {})

        q_text = sa.question_text
        if not q_text and sa.question_number in QUESTIONS_BY_ID:
            q_text = QUESTIONS_BY_ID[sa.question_number].get("question", "")

        answers_out.append({
            "question_number": sa.question_number,
            "question": q_text or f"Question {sa.question_number}",
            "options": opts,
            "selected_answer": sa.selected_answer,
            "correct_answer": sa.correct_answer,
            "marks": sa.marks or 1,
            "is_correct": sa.is_correct
        })

    return {
        "student": student,
        "attempt": attempt,
        "answers": answers_out,
        "security_logs": security_logs
    }


# =========================================================================
# 10. ADMIN DASHBOARD & STUDENT RECORDS
# =========================================================================

@router.get("/admin/dashboard", response_model=schemas.AdminDashboardStats)
def get_admin_dashboard(admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    total_students = db.query(models.Student).count()
    total_attempts = db.query(models.Attempt).count()
    completed_tests = db.query(models.Attempt).filter(models.Attempt.status.in_(["completed", "timed_out"])).count()
    
    avg_score = db.query(func.avg(models.Attempt.score)).filter(models.Attempt.status.in_(["completed", "timed_out"])).scalar() or 0.0
    avg_pct = db.query(func.avg(models.Attempt.percentage)).filter(models.Attempt.status.in_(["completed", "timed_out"])).scalar() or 0.0
    
    total_departments = db.query(models.Department).count()
    total_subjects = db.query(models.Subject).count()
    total_exams = db.query(models.Exam).count()
    total_questions = db.query(models.Question).count()

    return {
        "total_students": total_students,
        "total_attempts": total_attempts,
        "completed_tests": completed_tests,
        "average_score": round(float(avg_score), 2),
        "average_percentage": round(float(avg_pct), 2),
        "total_departments": total_departments,
        "total_subjects": total_subjects,
        "total_exams": total_exams,
        "total_questions": total_questions
    }


@router.get("/admin/students", response_model=List[schemas.StudentSummary])
def get_admin_students(
    search: Optional[str] = Query(None),
    department: Optional[str] = Query(None),
    admin: str = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    query = db.query(models.Student)
    if search:
        s = f"%{search.strip()}%"
        query = query.filter(or_(models.Student.name.ilike(s), models.Student.enrollment_no.ilike(s)))
    if department:
        query = query.filter(models.Student.department == department)

    students = query.order_by(models.Student.created_at.desc()).all()
    summaries = []

    for s in students:
        latest_attempt = db.query(models.Attempt).filter(
            models.Attempt.student_id == s.id
        ).order_by(models.Attempt.started_at.desc()).first()

        summaries.append({
            "student_id": s.id,
            "enrollment_no": s.enrollment_no,
            "name": s.name,
            "department": s.department,
            "gender": s.gender,
            "program": s.program,
            "year": s.year,
            "semester": s.semester,
            "is_active": s.is_active,
            "score": latest_attempt.score if latest_attempt else None,
            "percentage": latest_attempt.percentage if latest_attempt else None,
            "attempt_date": (latest_attempt.completed_at or latest_attempt.submitted_at or latest_attempt.started_at) if latest_attempt else None,
            "status": latest_attempt.status if latest_attempt else "unattempted",
            "attempt_id": latest_attempt.id if latest_attempt else None,
            "exam_title": latest_attempt.exam_title_snapshot if latest_attempt else None,
            "tab_switch_count": latest_attempt.tab_switch_count if latest_attempt else 0,
            "fullscreen_exit_count": latest_attempt.fullscreen_exit_count if latest_attempt else 0
        })

    return summaries


@router.get("/admin/students/{student_id}", response_model=schemas.TestResultOut)
def get_admin_student_details(
    student_id: int,
    attempt_id: Optional[int] = Query(None),
    admin: str = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    student = db.query(models.Student).filter(models.Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found.")

    if attempt_id:
        attempt = db.query(models.Attempt).filter(
            models.Attempt.id == attempt_id,
            models.Attempt.student_id == student_id
        ).first()
    else:
        attempt = db.query(models.Attempt).filter(
            models.Attempt.student_id == student_id
        ).order_by(models.Attempt.started_at.desc()).first()

    if not attempt:
        raise HTTPException(status_code=404, detail="No assessment attempt found for this student.")

    return get_attempt_result(attempt.id, db)


@router.delete("/admin/attempts/{attempt_id}")
def delete_assessment_attempt(
    attempt_id: int,
    admin: str = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    """
    CRITICAL REQUIREMENT: DELETE ASSESSMENT ATTEMPT
    Safely delete attempt and its answers/logs, preserving student and questions.
    """
    attempt = db.query(models.Attempt).filter(models.Attempt.id == attempt_id).first()
    if not attempt:
        raise HTTPException(status_code=404, detail="Assessment attempt not found.")

    student_id = attempt.student_id
    db.query(models.StudentAnswer).filter(models.StudentAnswer.attempt_id == attempt_id).delete(synchronize_session=False)
    db.query(models.ExamSecurityLog).filter(models.ExamSecurityLog.attempt_id == attempt_id).delete(synchronize_session=False)
    db.delete(attempt)
    db.commit()

    return {
        "status": "success",
        "message": f"Assessment attempt #{attempt_id} deleted successfully.",
        "attempt_id": attempt_id,
        "student_id": student_id
    }


# =========================================================================
# 11. QUESTION BANK MANAGEMENT (ADMIN)
# =========================================================================

@router.get("/admin/questions", response_model=List[schemas.AdminQuestionOut])
def get_admin_questions(
    subject_id: Optional[int] = None,
    department_id: Optional[int] = None,
    program_id: Optional[int] = None,
    academic_year_id: Optional[int] = None,
    semester_id: Optional[int] = None,
    topic: Optional[str] = None,
    difficulty: Optional[str] = None,
    is_active: Optional[bool] = None,
    search: Optional[str] = None,
    admin: str = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    q = db.query(models.Question)
    if subject_id:
        q = q.filter(models.Question.subject_id == subject_id)
    if department_id:
        q = q.filter(models.Question.department_id == department_id)
    if program_id:
        q = q.filter(models.Question.program_id == program_id)
    if academic_year_id:
        q = q.filter(models.Question.academic_year_id == academic_year_id)
    if semester_id:
        q = q.filter(models.Question.semester_id == semester_id)
    if topic:
        q = q.filter(models.Question.topic == topic)
    if difficulty:
        q = q.filter(models.Question.difficulty == difficulty)
    if is_active is not None:
        q = q.filter(models.Question.is_active == is_active)
    if search:
        s = f"%{search.strip()}%"
        q = q.filter(models.Question.question_text.ilike(s))

    return q.order_by(models.Question.id.asc()).all()


@router.get("/admin/questions/{question_id}", response_model=schemas.AdminQuestionOut)
def get_admin_question(question_id: int, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    q = db.query(models.Question).filter(models.Question.id == question_id).first()
    if not q:
        raise HTTPException(status_code=404, detail="Question not found.")
    return q


@router.post("/admin/questions", response_model=schemas.AdminQuestionOut, status_code=status.HTTP_201_CREATED)
def create_question(q_in: schemas.QuestionCreate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    correct = q_in.correct_answer.strip().upper()
    if correct not in ["A", "B", "C", "D"]:
        raise HTTPException(status_code=400, detail="Correct answer must be A, B, C, or D.")

    dept_name = None
    if q_in.department_id:
        dept = db.query(models.Department).filter(models.Department.id == q_in.department_id).first()
        dept_name = dept.name if dept else None

    subj_name = None
    if q_in.subject_id:
        subj = db.query(models.Subject).filter(models.Subject.id == q_in.subject_id).first()
        subj_name = subj.name if subj else None

    prog_name = None
    if q_in.program_id:
        prog = db.query(models.Program).filter(models.Program.id == q_in.program_id).first()
        prog_name = prog.name if prog else None

    year_name = None
    if q_in.academic_year_id:
        ay = db.query(models.AcademicYear).filter(models.AcademicYear.id == q_in.academic_year_id).first()
        year_name = ay.name if ay else None

    sem_name = None
    if q_in.semester_id:
        sem = db.query(models.Semester).filter(models.Semester.id == q_in.semester_id).first()
        sem_name = sem.name if sem else None

    now = datetime.utcnow()
    q = models.Question(
        question_text=q_in.question_text.strip(),
        option_a=q_in.option_a.strip(),
        option_b=q_in.option_b.strip(),
        option_c=q_in.option_c.strip(),
        option_d=q_in.option_d.strip(),
        correct_answer=correct,
        topic=q_in.topic.strip() if q_in.topic else None,
        department_id=q_in.department_id,
        department_name=dept_name,
        program_id=q_in.program_id,
        program_name=prog_name,
        academic_year_id=q_in.academic_year_id,
        year_name=year_name,
        semester_id=q_in.semester_id,
        semester_name=sem_name,
        subject_id=q_in.subject_id,
        subject_name=subj_name,
        difficulty=q_in.difficulty or "Medium",
        marks=q_in.marks or 1,
        is_active=q_in.is_active,
        created_at=now,
        updated_at=now
    )
    db.add(q)
    db.commit()
    db.refresh(q)
    return q


@router.put("/admin/questions/{question_id}", response_model=schemas.AdminQuestionOut)
def update_question(question_id: int, q_in: schemas.QuestionUpdate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    q = db.query(models.Question).filter(models.Question.id == question_id).first()
    if not q:
        raise HTTPException(status_code=404, detail="Question not found.")

    if q_in.correct_answer is not None:
        c = q_in.correct_answer.strip().upper()
        if c not in ["A", "B", "C", "D"]:
            raise HTTPException(status_code=400, detail="Correct answer must be A, B, C, or D.")
        q.correct_answer = c

    if q_in.question_text is not None:
        q.question_text = q_in.question_text.strip()
    if q_in.option_a is not None:
        q.option_a = q_in.option_a.strip()
    if q_in.option_b is not None:
        q.option_b = q_in.option_b.strip()
    if q_in.option_c is not None:
        q.option_c = q_in.option_c.strip()
    if q_in.option_d is not None:
        q.option_d = q_in.option_d.strip()
    if q_in.topic is not None:
        q.topic = q_in.topic.strip()
    if q_in.difficulty is not None:
        q.difficulty = q_in.difficulty
    if q_in.marks is not None:
        q.marks = q_in.marks
    if q_in.is_active is not None:
        q.is_active = q_in.is_active

    if q_in.subject_id is not None:
        subj = db.query(models.Subject).filter(models.Subject.id == q_in.subject_id).first()
        q.subject_id = subj.id if subj else None
        q.subject_name = subj.name if subj else None
    if q_in.department_id is not None:
        dept = db.query(models.Department).filter(models.Department.id == q_in.department_id).first()
        q.department_id = dept.id if dept else None
        q.department_name = dept.name if dept else None
    if q_in.program_id is not None:
        prog = db.query(models.Program).filter(models.Program.id == q_in.program_id).first()
        q.program_id = prog.id if prog else None
        q.program_name = prog.name if prog else None
    if q_in.academic_year_id is not None:
        ay = db.query(models.AcademicYear).filter(models.AcademicYear.id == q_in.academic_year_id).first()
        q.academic_year_id = ay.id if ay else None
        q.year_name = ay.name if ay else None
    if q_in.semester_id is not None:
        sem = db.query(models.Semester).filter(models.Semester.id == q_in.semester_id).first()
        q.semester_id = sem.id if sem else None
        q.semester_name = sem.name if sem else None

    q.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(q)
    return q


@router.patch("/admin/questions/{question_id}/status", response_model=schemas.AdminQuestionOut)
def toggle_question_status(question_id: int, status_in: schemas.QuestionStatusUpdate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    q = db.query(models.Question).filter(models.Question.id == question_id).first()
    if not q:
        raise HTTPException(status_code=404, detail="Question not found.")
    q.is_active = status_in.is_active
    q.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(q)
    return q


@router.post("/admin/questions/{question_id}/deactivate", response_model=schemas.AdminQuestionOut)
def deactivate_question(question_id: int, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    q = db.query(models.Question).filter(models.Question.id == question_id).first()
    if not q:
        raise HTTPException(status_code=404, detail="Question not found.")
    q.is_active = False
    q.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(q)
    return q


@router.delete("/admin/questions/{question_id}")
def delete_question(question_id: int, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    """
    CRITICAL REQUIREMENT: QUESTION DELETE
    If referenced by attempts, block hard deletion and recommend deactivate.
    If unused, permanently delete.
    """
    q = db.query(models.Question).filter(models.Question.id == question_id).first()
    if not q:
        raise HTTPException(status_code=404, detail="Question not found.")

    usage_count = db.query(models.StudentAnswer).filter(models.StudentAnswer.question_id == question_id).count()
    exam_usage = db.query(models.ExamQuestion).filter(models.ExamQuestion.question_id == question_id).count()

    if usage_count > 0 or exam_usage > 0:
        raise HTTPException(
            status_code=400,
            detail=f"Question #{question_id} cannot be permanently deleted because it is referenced in {usage_count} student answer(s) and {exam_usage} exam(s). Please deactivate it instead to protect historical assessment records."
        )

    db.delete(q)
    db.commit()
    return {"message": f"Question #{question_id} deleted permanently.", "id": question_id}


# =========================================================================
# 12. ASSESSMENT SETTINGS (ADMIN)
# =========================================================================

@router.get("/admin/assessment-settings", response_model=schemas.AssessmentSettingsOut)
def get_assessment_settings(admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    setting = get_or_create_settings(db)
    active_cnt = db.query(models.Question).filter(models.Question.is_active == True).count()
    total_cnt = db.query(models.Question).count()
    return {
        "question_count": setting.question_count,
        "time_limit_minutes": setting.time_limit_minutes,
        "active_questions_count": active_cnt,
        "total_questions_count": total_cnt,
        "updated_at": setting.updated_at
    }


@router.put("/admin/assessment-settings", response_model=schemas.AssessmentSettingsOut)
def update_assessment_settings(settings_in: schemas.AssessmentSettingsUpdate, admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    active_cnt = db.query(models.Question).filter(models.Question.is_active == True).count()
    if settings_in.question_count > active_cnt and active_cnt > 0:
        raise HTTPException(
            status_code=400,
            detail=f"Cannot configure {settings_in.question_count} questions: only {active_cnt} active questions exist."
        )

    setting = get_or_create_settings(db)
    setting.question_count = settings_in.question_count
    setting.time_limit_minutes = settings_in.time_limit_minutes
    setting.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(setting)

    total_cnt = db.query(models.Question).count()
    return {
        "question_count": setting.question_count,
        "time_limit_minutes": setting.time_limit_minutes,
        "active_questions_count": active_cnt,
        "total_questions_count": total_cnt,
        "updated_at": setting.updated_at
    }


# =========================================================================
# 13. EXPORTS (EXCEL & CSV)
# =========================================================================

@router.get("/admin/export/excel")
def export_students_excel(admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    students = db.query(models.Student).order_by(models.Student.enrollment_no.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assessment Records"

    headers = [
        "Student ID", "Enrollment No", "Student Name", "Department", "Gender",
        "Program", "Year", "Semester", "Exam Title", "Subject",
        "Score", "Percentage (%)", "Status", "Attempt Date",
        "Tab Switches", "Fullscreen Exits", "Copy Attempts", "Paste Attempts"
    ]
    ws.append(headers)

    for s in students:
        attempts = db.query(models.Attempt).filter(models.Attempt.student_id == s.id).order_by(models.Attempt.started_at.desc()).all()
        if not attempts:
            ws.append([
                s.id, s.enrollment_no, s.name, s.department, s.gender or "",
                s.program or "", s.year or "", s.semester or "", "N/A", "N/A",
                "N/A", "N/A", "Unattempted", "N/A",
                0, 0, 0, 0
            ])
        else:
            for att in attempts:
                att_date = att.completed_at or att.submitted_at or att.started_at
                ws.append([
                    s.id, s.enrollment_no, s.name, s.department, s.gender or "",
                    s.program or "", s.year or "", s.semester or "",
                    att.exam_title_snapshot or "General Assessment",
                    att.subject_name_snapshot or "DSA",
                    att.score, att.percentage, att.status,
                    att_date.strftime("%Y-%m-%d %H:%M:%S") if att_date else "N/A",
                    att.tab_switch_count or 0,
                    att.fullscreen_exit_count or 0,
                    att.copy_count or 0,
                    att.paste_count or 0
                ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"college_assessment_records_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/admin/export/csv")
def export_students_csv(admin: str = Depends(get_current_admin), db: Session = Depends(get_db)):
    students = db.query(models.Student).order_by(models.Student.enrollment_no.asc()).all()

    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow([
        "Student ID", "Enrollment No", "Student Name", "Department", "Gender",
        "Program", "Year", "Semester", "Exam Title", "Subject",
        "Score", "Percentage", "Status", "Attempt Date",
        "Tab Switches", "Fullscreen Exits", "Copy Attempts", "Paste Attempts"
    ])

    for s in students:
        attempts = db.query(models.Attempt).filter(models.Attempt.student_id == s.id).order_by(models.Attempt.started_at.desc()).all()
        if not attempts:
            writer.writerow([
                s.id, s.enrollment_no, s.name, s.department, s.gender or "",
                s.program or "", s.year or "", s.semester or "", "N/A", "N/A",
                "N/A", "N/A", "Unattempted", "N/A",
                0, 0, 0, 0
            ])
        else:
            for att in attempts:
                att_date = att.completed_at or att.submitted_at or att.started_at
                writer.writerow([
                    s.id, s.enrollment_no, s.name, s.department, s.gender or "",
                    s.program or "", s.year or "", s.semester or "",
                    att.exam_title_snapshot or "General Assessment",
                    att.subject_name_snapshot or "DSA",
                    att.score, att.percentage, att.status,
                    att_date.strftime("%Y-%m-%d %H:%M:%S") if att_date else "N/A",
                    att.tab_switch_count or 0,
                    att.fullscreen_exit_count or 0,
                    att.copy_count or 0,
                    att.paste_count or 0
                ])

    output.seek(0)
    filename = f"college_assessment_records_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.csv"
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8")),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
